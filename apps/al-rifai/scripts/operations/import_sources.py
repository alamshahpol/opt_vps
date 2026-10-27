"""Read-only historical extraction. Never connects to or writes Fazle Core."""
import re,json,hashlib,sys,glob
from pathlib import Path
from datetime import datetime
from decimal import Decimal
PHONE=r'(?:\+?880[\s-]?1[3-9](?:[\s-]?\d){8}|0?1[3-9](?:[\s-]?\d){8})'
PAY=re.compile(rf'^\s*(?:ID\s*[:\-]?\s*(?P<eid>{PHONE})\s+)?(?P<name>.*?)\s*(?P<phone>{PHONE})\s*\(\s*(?P<method>Nagad|N|Bikash|Bkash|B|Cash|C|Rocket|R|BANK)\s*\)\s*(?P<amount>[\d,]+(?:\.\d{{1,2}})?)\s*/[-+]?',re.I)
METHOD={'b':'bkash','bkash':'bkash','bikash':'bkash','n':'nagad','nagad':'nagad','c':'cash','cash':'cash','r':'rocket','rocket':'rocket','bank':'bank'}
def phone(v):
 if isinstance(v,(int,float)):v=str(int(v))
 n=re.sub(r'\D','',str(v or ''))
 if n.startswith('880'): n='0'+n[3:]
 if len(n)==10:n='0'+n
 return n if re.fullmatch(r'01[3-9]\d{8}',n) else ''
def ident(s):return hashlib.sha256(s.encode()).hexdigest()[:28]
def parse_line(s):
 m=PAY.match(s)
 if not m:return None
 v=m.groupdict();p=phone(v['phone']);eid=phone(v['eid'])
 amount=int(Decimal(v['amount'].replace(',',''))*100)
 if not p or amount<=0:return None
 name=v['name'].strip(' ,;:-');issues=[]
 if re.search(PHONE,name):return None
 if not eid:issues.append('Employee identity inferred from payout number')
 if not name:issues.append('Employee name missing')
 if '<This message was edited>' in s:issues.append('Edited source message')
 if s[m.end():].strip() not in ('','-','<This message was edited>'):issues.append('Additional text needs review')
 return dict(employeePhone=eid or p,payoutPhone=p,name=name or 'Name unavailable',amountCents=amount,method=METHOD[v['method'].lower()],issues=issues)
def extract(path):
 text=Path(path).read_text();messages=[]
 for line_no,line in enumerate(text.splitlines(),1):
  m=re.match(r'^(\d{1,2}/\d{1,2}/\d{4}), (\d{2}:\d{2}) - ([^:]+): ?(.*)$',line)
  if m:
   d,t,s,b=m.groups();messages.append(dict(id=ident('mamun:'+str(line_no)+':'+line),date=datetime.strptime(d,'%d/%m/%Y').date().isoformat(),time=t,sender=s,body=b,line=line_no))
  elif messages:messages[-1]['body']+='\n'+line
 rows=[];summaries=[];reviews=[]
 for m in messages:
  if m['sender']=='Mamun Vai Bkash Nagod':
   if re.search(r'জমা|বাকি|টোটাল|total|due|balance',m['body'],re.I):
    d=re.search(r'\b(\d{1,2}/\d{1,2}/\d{2,4})\s*=',m['body']);day=m['date']
    if d:
     try:day=datetime.strptime(d[1],'%d/%m/%y' if len(d[1].split('/')[-1])==2 else '%d/%m/%Y').date().isoformat()
     except ValueError:pass
    summaries.append({**m,'accountingDate':day})
   continue
  if m['sender']!='Muradul Azim':continue
  found=0
  for i,line in enumerate(m['body'].splitlines()):
   p=parse_line(line)
   if p:
    if '<This message was edited>' in m['body'] and 'Edited source message' not in p['issues']:p['issues'].append('Edited source message')
    found+=1;rows.append({**p,'id':ident(m['id']+':'+str(i)),'date':m['date'],'time':m['time'],'source':'WhatsApp archive','sourceId':m['id'],'sourceLine':m['line'],'sourceText':m['body'],'sender':m['sender'],'category':'Unclassified payment','status':'Historical instruction','purpose':''})
   elif re.search(PHONE,line) and re.search(r'\(\s*[BNCR]\s*\)|\d\s*/[-+]',line,re.I):reviews.append({**m,'id':ident(m['id']+':review:'+str(i)),'reason':'Payment-like line does not match the required format'})
 return rows,summaries,reviews,messages
if __name__=='__main__':
 import openpyxl
 output=Path(sys.argv[2]).resolve()
 evidence=Path('/var/lib/vps/al-rifai/evidence').resolve()
 if output==evidence or evidence in output.parents:raise SystemExit('Runtime evidence is immutable')
 root=Path(sys.argv[1]);out=Path(sys.argv[2]);rows,sums,revs,msgs=extract(root/'upload/01-WhatsApp-Chat-with-Mamun-Vai-Bkash-Nagod.txt')
 excel=[];programs=[];bills=[];audit=[]
 for file in glob.glob(str(root/'attachments/*/*.xlsx')):
  w=openpyxl.load_workbook(file,data_only=True);wf=openpyxl.load_workbook(file,data_only=False)
  audit.append({'file':Path(file).name,'sheets':[{'name':s.title,'rows':s.max_row,'formulas':sum(c.data_type=='f' for row in wf[s.title] for c in row)} for s in w]})
  if 'Cashbook' in file:
   for i,r in enumerate(w.active.values,1):
    if i==1 or not r[0] or not isinstance(r[5],(int,float)):continue
    dt=datetime.strptime(str(r[0]),'%d.%m.%Y').date().isoformat();eid=phone(r[1]);payout=phone(r[3]);remarks=str(r[4] or '')
    method=METHOD.get(remarks.lower(),'unknown');issues=[]
    if not eid:issues.append('Employee ID missing')
    if not payout:issues.append('Payout number missing')
    if method=='unknown':issues.append('Method and purpose require separation')
    excel.append(dict(id=ident('cashbook:MasterBook:'+str(i)),employeePhone=eid,payoutPhone=payout,name=str(r[2] or 'Name unavailable'),amountCents=int(Decimal(str(r[5]))*100),date=dt,time='',source='Excel Cashbook',sourceId='MasterBook:'+str(i),sourceLine=i,sourceText=json.dumps(list(r[:6]),ensure_ascii=False),sender='Cashbook',category='Unclassified payment',status='Historical workbook',purpose=remarks,method=method,issues=issues))
  elif 'Escort Program' in file:
   def dt(v):
    if not v:return ('','')
    m=re.search(r'(\d{1,2}\.\d{1,2}\.\d{4})',str(v));sh=re.search(r'\(([DN])\)',str(v))
    return (datetime.strptime(m[1],'%d.%m.%Y').date().isoformat() if m else '',sh[1] if sh else '')
   for i,r in enumerate(w['SanMarineTCIS'].values,1):
    if i==1 or not r[0]:continue
    start,shift=dt(r[3]);end,eshift=dt(r[4]);point=str(r[7] or '');status='Cancelled' if point=='Cancel' else 'Escort change' if point=='Escort Change' else 'Released' if end else 'Recorded'
    programs.append(dict(id=ident('escort-workbook:'+str(i)),motherVessel=str(r[0]),lighterVessel=str(r[1] or ''),escortName=str(r[2] or ''),escortMobile=phone(r[10]),masterMobile='',destination='' if point in ('Cancel','Escort Change','Ship unloaded','—') else point,startDate=start,endDate=end,shift=shift,status=status,source='Escort workbook · SanMarineTCIS',notes='Historical workbook row '+str(i),duplicateOf=None,createdAt='2026-10-02T00:00:00Z',revision=0,details={'recordKind':'Program','sourceLine':str(i),'issues':['Historical workbook: confirm current status'],'releasePoint':point,'endShift':eshift}))
   mother=''
   for i,r in enumerate(w['BillGenerate '].values,1):
    if isinstance(r[1],str) and r[1].startswith('Mother Vessel:'):mother=r[1].split(':',1)[1].strip()
    if isinstance(r[1],(int,float)) and isinstance(r[6],(int,float)) and r[2] and dt(r[4])[0] and dt(r[5])[0]:
     bills.append(dict(id='bill:'+str(i),motherVessel=mother,lighterVessel=str(r[2]),escortName=str(r[3]),startDate=str(r[4]),endDate=str(r[5]),days=r[6],serviceCharge=r[7],releasePoint=str(r[8] or ''),conveyance=r[9] or 0,total=r[10],sourceLine=i))
 data={'cash':rows,'excel':excel,'summaries':sums,'reviews':revs,'programs':programs,'bills':bills,'analysis':{'messageCount':len(msgs),'firstDate':msgs[0]['date'],'lastDate':msgs[-1]['date'],'chatRows':len(rows),'excelRows':len(excel),'summaryMessages':len(sums),'unparsedPayments':len(revs),'editedMessages':sum('<This message was edited>' in m['body'] for m in msgs),'missingMedia':sum('<Media omitted>' in m['body'] for m in msgs),'workbooks':audit,'repositoryCommit':'d5551931a53b7b10bd89072ac1b48c72a4fcf8d9'}}
 out.open('x').write(json.dumps(data,ensure_ascii=False,separators=(',',':')));print(json.dumps(data['analysis'],ensure_ascii=False));print('programs',len(programs),'bill details',len(bills))
